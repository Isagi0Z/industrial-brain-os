"""Unit tests for M3 — OCR & Layout-Aware Parsing.

All external dependencies (storage, DB, Redis) are mocked.
"""

from __future__ import annotations

from datetime import datetime, timezone
from unittest.mock import MagicMock
import io
import uuid

import pytest

from app.domain.document.constants import ChunkType, JobStatus, DocumentStatus
from app.domain.document.models import (
    Document,
    DocumentChunk,
    DocumentVersion,
)
from app.infrastructure.document.parsing.chunker import (
    count_tokens,
    finalize_chunks,
    _split_text_with_overlap,
)


# ---------------------------------------------------------------------------
# Fixtures
# ---------------------------------------------------------------------------


def _now():
    return datetime.now(timezone.utc)


def _make_doc(doc_id="doc-1", mime="application/pdf"):
    return Document(
        id=doc_id,
        original_filename="test.pdf",
        mime_type=mime,
        size_bytes=100,
        sha256_hash="abc123",
        status=DocumentStatus.READY_FOR_PROCESSING,
        created_at=_now(),
        updated_at=_now(),
        created_by="user-1",
    )


def _make_version(doc_id="doc-1"):
    return DocumentVersion(
        id=str(uuid.uuid4()),
        document_id=doc_id,
        version_number=1,
        storage_path=f"{doc_id}/v1/test.pdf",
        size_bytes=100,
        created_at=_now(),
        created_by="user-1",
    )


def _make_chunk(doc_id="doc-1", chunk_type=ChunkType.PARAGRAPH, text="hello world"):
    return DocumentChunk(
        id=str(uuid.uuid4()),
        document_id=doc_id,
        chunk_index=0,
        chunk_type=chunk_type,
        text=text,
        page_number=1,
        parent_section_header=None,
        bbox_json=None,
        token_count=count_tokens(text),
        created_at=_now(),
    )


# ---------------------------------------------------------------------------
# chunker unit tests
# ---------------------------------------------------------------------------


def test_count_tokens_basic():
    assert count_tokens("hello world") > 0


def test_finalize_chunks_passthrough_short():
    chunk = _make_chunk(text="Short paragraph text.")
    result = finalize_chunks([chunk])
    assert len(result) == 1
    assert result[0].chunk_type == ChunkType.PARAGRAPH


def test_finalize_chunks_table_not_split():
    # A very long table should remain a single chunk
    long_text = " | ".join([f"cell {i}" for i in range(300)])
    chunk = _make_chunk(chunk_type=ChunkType.TABLE, text=long_text)
    result = finalize_chunks([chunk])
    assert len(result) == 1
    assert result[0].chunk_type == ChunkType.TABLE


def test_finalize_chunks_long_paragraph_split():
    long_text = "word " * 1000
    chunk = _make_chunk(text=long_text)
    result = finalize_chunks([chunk])
    assert len(result) > 1
    for i, c in enumerate(result):
        assert c.chunk_index == i


def test_split_with_overlap_short_text():
    text = "short text"
    parts = _split_text_with_overlap(text, soft_max=512, overlap=64)
    assert parts == [text]


def test_split_with_overlap_produces_overlap():
    # Need to exceed CHUNK_HARD_MAX_TOKENS (768) to trigger splitting
    long_text = "word " * 1000
    parts = _split_text_with_overlap(long_text)
    assert len(parts) > 1


def test_finalize_chunks_reindexes_sequentially():
    chunks = [_make_chunk(text=f"Paragraph {i}.") for i in range(5)]
    result = finalize_chunks(chunks)
    assert [c.chunk_index for c in result] == list(range(len(result)))


# ---------------------------------------------------------------------------
# DocumentParsingUseCase unit tests
# ---------------------------------------------------------------------------


@pytest.fixture()
def mock_parsing_use_case():
    from app.application.document.parsing_service import DocumentParsingUseCase

    doc_repo = MagicMock()
    storage = MagicMock()
    job_repo = MagicMock()
    chunk_repo = MagicMock()
    parser = MagicMock()

    uc = DocumentParsingUseCase(
        document_repo=doc_repo,
        storage_service=storage,
        job_repo=job_repo,
        chunk_repo=chunk_repo,
        parsers=[parser],
    )
    return uc, doc_repo, storage, job_repo, chunk_repo, parser


def test_parse_document_success(mock_parsing_use_case):
    uc, doc_repo, storage, job_repo, chunk_repo, parser = mock_parsing_use_case

    doc = _make_doc()
    version = _make_version()
    chunk = _make_chunk()

    doc_repo.get_by_id.return_value = doc
    doc_repo.get_versions.return_value = [version]
    storage.download_file.return_value = b"%PDF-1.4 test"
    parser.can_parse.return_value = True
    parser.parse.return_value = [chunk]

    result = uc.parse_document("doc-1", "job-1")

    assert len(result) == 1
    job_repo.update_status.assert_any_call("job-1", JobStatus.EXTRACTING)
    job_repo.update_status.assert_any_call("job-1", JobStatus.EXTRACTED)
    job_repo.update_status.assert_any_call("job-1", JobStatus.PARSING)
    job_repo.update_status.assert_any_call("job-1", JobStatus.CHUNKED)
    chunk_repo.bulk_insert.assert_called_once()


def test_parse_document_missing_doc(mock_parsing_use_case):
    uc, doc_repo, storage, job_repo, chunk_repo, parser = mock_parsing_use_case
    doc_repo.get_by_id.return_value = None

    result = uc.parse_document("missing-doc", "job-2")

    assert result == []
    job_repo.update_status.assert_any_call(
        "job-2", JobStatus.FAILED, error_message="Document missing-doc not found."
    )


def test_parse_document_no_version(mock_parsing_use_case):
    uc, doc_repo, storage, job_repo, chunk_repo, parser = mock_parsing_use_case
    doc_repo.get_by_id.return_value = _make_doc()
    doc_repo.get_versions.return_value = []

    result = uc.parse_document("doc-1", "job-3")

    assert result == []
    calls = [str(c) for c in job_repo.update_status.call_args_list]
    assert any("FAILED" in c for c in calls)


def test_parse_document_storage_failure(mock_parsing_use_case):
    uc, doc_repo, storage, job_repo, chunk_repo, parser = mock_parsing_use_case
    doc_repo.get_by_id.return_value = _make_doc()
    doc_repo.get_versions.return_value = [_make_version()]
    storage.download_file.side_effect = Exception("MinIO connection refused")

    result = uc.parse_document("doc-1", "job-4")

    assert result == []
    calls = [str(c) for c in job_repo.update_status.call_args_list]
    assert any("FAILED" in c for c in calls)


def test_parse_document_no_parser_for_mime(mock_parsing_use_case):
    uc, doc_repo, storage, job_repo, chunk_repo, parser = mock_parsing_use_case
    doc = _make_doc(mime="application/x-unknown")
    doc_repo.get_by_id.return_value = doc
    doc_repo.get_versions.return_value = [_make_version()]
    storage.download_file.return_value = b"data"
    parser.can_parse.return_value = False

    result = uc.parse_document("doc-1", "job-5")

    assert result == []
    calls = [str(c) for c in job_repo.update_status.call_args_list]
    assert any("FAILED" in c for c in calls)


# ---------------------------------------------------------------------------
# PyMuPDF parser smoke test (no live MinIO required)
# ---------------------------------------------------------------------------


def test_pymupdf_parser_can_parse():
    from app.infrastructure.document.parsing.pymupdf_parser import PyMuPDFParser

    p = PyMuPDFParser()
    assert p.can_parse("application/pdf")
    assert not p.can_parse("image/png")


def test_pymupdf_parser_produces_chunks():
    import fitz

    from app.infrastructure.document.parsing.pymupdf_parser import PyMuPDFParser

    # Minimal valid PDF in memory
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((72, 100), "Section 1: Introduction", fontsize=18)
    page.insert_text(
        (72, 130),
        "This is a paragraph of industrial text about pumps and valves.",
        fontsize=11,
    )
    buf = io.BytesIO()
    doc.save(buf)
    doc.close()
    pdf_bytes = buf.getvalue()

    parser = PyMuPDFParser()
    chunks = parser.parse(
        document_id="doc-test",
        filename="test.pdf",
        mime_type="application/pdf",
        file_bytes=pdf_bytes,
    )
    assert len(chunks) >= 1
    for c in chunks:
        assert c.document_id == "doc-test"
        assert c.chunk_type in ChunkType.__members__.values()
        assert c.token_count > 0


# ---------------------------------------------------------------------------
# DocxParser smoke test
# ---------------------------------------------------------------------------


def test_docx_parser_can_parse():
    from app.infrastructure.document.parsing.docx_parser import DocxParser

    p = DocxParser()
    assert p.can_parse(
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    )
    assert not p.can_parse("application/pdf")


def test_docx_parser_produces_chunks():
    from docx import Document as DocxDoc
    from app.infrastructure.document.parsing.docx_parser import DocxParser

    doc = DocxDoc()
    doc.add_heading("M3 Test Document", level=1)
    doc.add_paragraph("This is a test paragraph about industrial equipment.")
    buf = io.BytesIO()
    doc.save(buf)
    docx_bytes = buf.getvalue()

    parser = DocxParser()
    chunks = parser.parse(
        document_id="doc-docx",
        filename="test.docx",
        mime_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        file_bytes=docx_bytes,
    )
    assert len(chunks) >= 1
    types = {c.chunk_type for c in chunks}
    assert ChunkType.HEADING in types or ChunkType.PARAGRAPH in types
