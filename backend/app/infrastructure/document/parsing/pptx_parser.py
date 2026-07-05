"""PPTX parser using python-pptx. One chunk stream per presentation, with slide
titles as headings and slide index as the page number (preserves hierarchy)."""

from __future__ import annotations

import io
import logging
import uuid
from datetime import datetime, timezone
from typing import List, Optional

from app.domain.document.constants import ChunkType
from app.domain.document.interfaces import IDocumentParser
from app.domain.document.models import DocumentChunk
from app.infrastructure.document.parsing.chunker import count_tokens, finalize_chunks

logger = logging.getLogger(__name__)

try:
    from pptx import Presentation  # type: ignore

    _PPTX_AVAILABLE = True
except Exception:  # noqa: BLE001
    Presentation = None  # type: ignore
    _PPTX_AVAILABLE = False
    logger.warning("python-pptx unavailable — PPTX files degrade to a filename note.")

_MIMES = {
    "application/vnd.openxmlformats-officedocument.presentationml.presentation",
    "application/vnd.ms-powerpoint",
}


class PptxParser(IDocumentParser):
    def can_parse(self, mime_type: str) -> bool:
        return (mime_type or "").split(";")[0].strip() in _MIMES

    def parse(
        self,
        document_id: str,
        filename: str,
        mime_type: str,
        file_bytes: bytes,
    ) -> List[DocumentChunk]:
        now = datetime.now(timezone.utc)
        if not _PPTX_AVAILABLE or Presentation is None:
            return [
                DocumentChunk(
                    id=str(uuid.uuid4()),
                    document_id=document_id,
                    chunk_index=0,
                    chunk_type=ChunkType.PARAGRAPH,
                    text=f"[Presentation: {filename} — python-pptx not installed]",
                    page_number=1,
                    parent_section_header=None,
                    bbox_json=None,
                    token_count=8,
                    created_at=now,
                )
            ]

        prs = Presentation(io.BytesIO(file_bytes))
        raw: List[DocumentChunk] = []
        idx = 0
        for slide_no, slide in enumerate(prs.slides, start=1):
            title: Optional[str] = None
            try:
                if slide.shapes.title and slide.shapes.title.text.strip():
                    title = slide.shapes.title.text.strip()
                    raw.append(
                        self._mk(
                            document_id,
                            idx,
                            ChunkType.HEADING,
                            title,
                            slide_no,
                            None,
                            now,
                        )
                    )
                    idx += 1
            except Exception:  # noqa: BLE001
                title = None

            for shape in slide.shapes:
                if shape.has_table:
                    table_data = []
                    cells = []
                    for r_idx, row in enumerate(shape.table.rows):
                        for c_idx, cell in enumerate(row.cells):
                            v = cell.text.strip()
                            if v:
                                table_data.append(
                                    {"row": r_idx, "col": c_idx, "value": v}
                                )
                                cells.append(v)
                    if cells:
                        raw.append(
                            self._mk(
                                document_id,
                                idx,
                                ChunkType.TABLE,
                                " | ".join(cells),
                                slide_no,
                                title,
                                now,
                                table_data,
                            )
                        )
                        idx += 1
                    continue
                if not shape.has_text_frame:
                    continue
                body = "\n".join(
                    p.text
                    for p in shape.text_frame.paragraphs
                    if p.text and p.text.strip()
                ).strip()
                if body and body != title:
                    raw.append(
                        self._mk(
                            document_id,
                            idx,
                            ChunkType.PARAGRAPH,
                            body,
                            slide_no,
                            title,
                            now,
                        )
                    )
                    idx += 1
        return finalize_chunks(raw)

    @staticmethod
    def _mk(
        document_id, idx, ctype, text, page, header, now, table_data=None
    ) -> DocumentChunk:
        return DocumentChunk(
            id=str(uuid.uuid4()),
            document_id=document_id,
            chunk_index=idx,
            chunk_type=ctype,
            text=text,
            page_number=page,
            parent_section_header=header,
            bbox_json=None,
            token_count=count_tokens(text),
            created_at=now,
            table_data_json=table_data,
        )
